from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\Works_Updates\NetSolution\src\SOFCMeas")
OUT = Path(r"C:\SOFCMeas\SOFCMeas 소프트웨어 운영 매뉴얼.docx")
LOGO = ROOT / "SOFCMeas" / "AMOCI.jpg"
MAIN_SCREEN = ROOT / "artifacts" / "verify-start-gate" / "StationPreview.png"
REVIEW_SCREEN = ROOT / "artifacts" / "verify-review-summary" / "ReviewSelection.png"
SETTING_SCREEN = ROOT / "artifacts" / "verify-full-settings" / "SettingPreview.png"

NAVY = "233247"
TEAL = "008A8E"
LIGHT_TEAL = "E8F3F3"
LIGHT_BLUE = "EAF1F6"
LIGHT_GRAY = "F3F5F7"
MID_GRAY = "C8D0D6"
DARK = "263238"
RED = "B43838"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, name="Malgun Gothic", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(paragraph, before=0, after=5, line=1.18):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_text(paragraph, text, bold=False, color=DARK, size=9.5):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return run


def add_body(doc, text, bold_lead=None, after=5):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=after)
    if bold_lead and text.startswith(bold_lead):
        add_text(p, bold_lead, bold=True)
        add_text(p, text[len(bold_lead):])
    else:
        add_text(p, text)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.left_indent = Cm(0.55 + 0.55 * level)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        set_paragraph_spacing(p, after=3, line=1.12)
        add_text(p, item, size=9.2)


def add_steps(doc, items):
    for idx, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.65)
        set_paragraph_spacing(p, after=4, line=1.14)
        add_text(p, f"{idx}. ", bold=True, color=TEAL, size=10)
        add_text(p, item, size=9.4)


def add_callout(doc, title, text, kind="note"):
    color = TEAL if kind == "note" else RED
    fill = LIGHT_TEAL if kind == "note" else "FBECEC"
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(17.2)
    shade(cell, fill)
    set_cell_margins(cell, 120, 180, 120, 180)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=2, line=1.12)
    add_text(p, title + "  ", bold=True, color=color, size=9.5)
    add_text(p, text, color=DARK, size=9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers, rows, widths=None, font_size=8.6):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, header in enumerate(headers):
        cell = hdr.cells[i]
        shade(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = Cm(widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, after=0, line=1.0)
        add_text(p, header, bold=True, color="FFFFFF", size=8.7)
        set_cell_margins(cell)
    for r_idx, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            if widths:
                cell.width = Cm(widths[i])
            if r_idx % 2 == 1:
                shade(cell, LIGHT_GRAY)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            set_paragraph_spacing(p, after=0, line=1.08)
            add_text(p, str(value), size=font_size)
            set_cell_margins(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    if getattr(doc, "_sofc_next_heading_new_page", False):
        p.paragraph_format.page_break_before = True
        doc._sofc_next_heading_new_page = False
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size=16 if level == 1 else 11.5, bold=True, color=NAVY if level == 1 else TEAL)
    return p


def add_image(doc, path, width, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=width)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(cap, after=6)
    add_text(cap, caption, color="5F6B73", size=8)


def page_break(doc):
    doc._sofc_next_heading_new_page = True


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    for name, size, color, before, after in (
        ("Heading 1", 16, NAVY, 8, 8),
        ("Heading 2", 11.5, TEAL, 7, 4),
        ("Heading 3", 10, DARK, 5, 3),
    ):
        style = styles[name]
        style.font.name = "Malgun Gothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(paragraph, "SOFCMeas 운영 매뉴얼  |  ", color="6B7780", size=8)
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def build_document():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)
    add_page_field(section.footer.paragraphs[0])

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(65)
    p.add_run().add_picture(str(LOGO), width=Inches(1.5))
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=8)
    add_text(p, "SOFCMeas", bold=True, color=TEAL, size=18)
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=6)
    add_text(p, "소프트웨어 운영 매뉴얼", bold=True, color=NAVY, size=29)
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=30)
    add_text(p, "설비 1 및 설비 2 하중 검사 운영 절차", color="52616B", size=13)
    line = doc.add_paragraph()
    line.paragraph_format.space_after = Pt(22)
    run = line.add_run("━" * 28)
    set_run_font(run, size=12, color=TEAL)
    add_table(doc, ["문서 구분", "내용"], [
        ["문서 버전", "1.0"],
        ["작성일", "2026-09-06"],
        ["대상 프로그램", "SOFCMeas"],
        ["적용 범위", "프로그램 실행, 검사, LOT 저장, 결과 조회, 로그 확인, 설정"],
    ], widths=[4.0, 13.0], font_size=9.2)
    add_callout(doc, "문서 범위", "이 문서는 소프트웨어 운영만 설명합니다. 광학계 조정과 기구부 설정은 포함하지 않습니다.")

    page_break(doc)
    add_heading(doc, "제1장  매뉴얼 사용 범위")
    add_body(doc, "SOFCMeas는 설비 1과 설비 2의 PLC에서 하중 데이터를 읽고, 검사 종료 시 판정한 뒤 LOT 단위 CSV 파일로 저장합니다. 두 설비는 같은 화면에서 독립적으로 동작하며 데이터와 상태도 각각 관리됩니다.")
    add_heading(doc, "운영 흐름", 2)
    add_table(doc, ["순서", "작업", "확인할 내용"], [
        ["1", "프로그램 실행", "사용 권한 확인 후 MAIN 화면 표시"],
        ["2", "PLC 상태 확인", "설비 1과 설비 2가 READY 또는 정상 연결 상태인지 확인"],
        ["3", "작업정보 입력", "모델명, LOT NO., 작업자, SPEC 확인"],
        ["4", "검사 대기", "START 버튼을 눌러 PLC START_REQ를 기다림"],
        ["5", "검사와 판정", "START_REQ에서 데이터 수집, END_REQ에서 최고점으로 GOOD 또는 NG 판정"],
        ["6", "LOT 저장", "LOT END를 눌러 누적 결과를 HDD에 저장"],
        ["7", "결과 확인", "날짜 조회 후 검사 번호를 선택하고 REVIEW"],
    ], widths=[1.3, 4.2, 11.5])
    add_heading(doc, "사용자 권한", 2)
    add_table(doc, ["권한", "사용 가능 기능"], [
        ["OPERATOR", "MAIN 화면에서 검사 실행, LOT 저장, 저장 결과 조회"],
        ["ADMINISTRATOR", "OPERATOR 기능과 LOG VIEW, SETTING 조회 및 저장"],
    ], widths=[4.0, 13.0])
    add_body(doc, "화면 상단의 USER 표시를 누르면 권한 창이 열립니다. 관리자 암호를 입력하면 ADMINISTRATOR로 전환됩니다. 설정 파일에 암호가 없을 때의 초기 암호는 1234입니다. 현장 적용 전에 반드시 변경하십시오.")

    page_break(doc)
    add_heading(doc, "제2장  프로그램 실행과 시작 점검")
    add_heading(doc, "프로그램 실행", 2)
    add_steps(doc, [
        "PC와 두 설비의 PLC 전원, 통신 케이블 연결 상태를 확인합니다.",
        "SOFCMeas 실행 파일 또는 현장에 등록된 바로가기를 실행합니다.",
        "시작 화면이 닫힌 뒤 MAIN 화면이 표시되는지 확인합니다.",
        "상단의 설비 1과 설비 2 연결 상태 및 각 설비 화면의 상태 표시를 확인합니다.",
    ])
    add_heading(doc, "사용 권한 검사", 2)
    add_body(doc, "프로그램은 실행할 때 PC의 물리적 네트워크 어댑터 주소인 MAC Address와 40일 사용 기간을 확인합니다. 등록된 MAC Address가 여러 개이면 그중 하나만 일치해도 실행됩니다.")
    add_table(doc, ["표시 메시지", "의미와 조치"], [
        ["You do not have permission to use this software.", "등록된 MAC Address가 없습니다. 프로그램이 종료되므로 관리자 또는 공급사에 등록을 요청합니다."],
        ["The software usage period has expired.", "40일 사용 기간이 끝났습니다. 관리자 또는 공급사에 기간 갱신 또는 영구 버전 적용을 요청합니다."],
        ["The software authorization could not be verified because the system clock was changed.", "PC 시간이 이전 실행 시각보다 뒤로 변경되었습니다. 날짜와 시간을 정상화한 뒤 관리자에게 확인을 요청합니다."],
        ["The software authorization could not be verified.", "권한 정보 확인에 실패했습니다. 임의로 레지스트리를 수정하지 말고 관리자에게 문의합니다."],
    ], widths=[7.0, 10.0], font_size=8.2)
    add_callout(doc, "주의", "PC 날짜, 시간, 권한 레지스트리를 임의로 변경하면 프로그램이 종료될 수 있습니다.", "warning")
    add_heading(doc, "일일 시작 점검", 2)
    add_bullets(doc, [
        "Windows 날짜와 시간이 현재 시각과 일치하는지 확인합니다.",
        "C:\\SOFCMeas 드라이브의 여유 공간과 Data, Log 폴더 접근 여부를 확인합니다.",
        "설비별 PLC IP가 현장 구성과 일치하는지 확인합니다. 기본값은 설비 1 172.20.9.100, 설비 2 172.20.9.101입니다.",
        "이전 작업의 미저장 데이터가 남아 있지 않은지 확인합니다.",
    ])

    page_break(doc)
    add_heading(doc, "제3장  MAIN 화면 구성")
    add_image(doc, MAIN_SCREEN, Inches(5.15), "그림 1  설비별 MAIN 화면 예시")
    add_table(doc, ["영역", "기능"], [
        ["설비 제목과 상태", "설비 번호, PLC IP, READY 또는 연결 상태를 표시합니다."],
        ["작업정보 입력", "모델명, 작업자, LOT NO., SPEC을 확인하거나 입력합니다."],
        ["실시간 하중 그래프", "X축은 읽기 횟수, Y축은 하중 kgf입니다. 검사 종료 후 최고점을 표시합니다."],
        ["검사현황", "검사수량, GOOD, NG, Yield(%), 현재 판정을 표시합니다."],
        ["LOG", "해당 설비의 연결, 수집, 판정, 저장 상태를 시간순으로 표시합니다."],
        ["START 또는 STOP", "START는 수집 대기를 시작합니다. 대기 중에는 STOP으로 바뀌며 누르면 수집을 중단합니다."],
        ["LOT END", "메모리에 누적된 완료 검사 데이터를 CSV 파일로 저장합니다."],
        ["저장 결과 조회", "날짜 범위로 저장 파일을 찾고 검사 번호별 그래프를 REVIEW합니다."],
    ], widths=[4.2, 12.8], font_size=7.8)

    page_break(doc)
    add_heading(doc, "제4장  검사 준비")
    add_heading(doc, "작업정보 입력", 2)
    add_steps(doc, [
        "검사할 설비 화면에서 작업정보 입력 영역을 선택합니다.",
        "모델명, LOT NO., 작업자를 입력합니다. 모델명과 LOT NO.는 저장 파일명에도 사용됩니다.",
        "SPEC 값을 확인합니다. 판정은 검사 최고점이 SPEC 하한값 이상이면 GOOD, 미만이면 NG입니다.",
        "입력한 정보가 현재 생산 지시와 일치하는지 다시 확인합니다.",
    ])
    add_callout(doc, "권장", "작업정보를 입력하지 않고 START_REQ가 들어오면 프로그램이 임시값을 만들 수 있습니다. 모델명은 STATION-N_날짜시간, LOT NO.는 UNASSIGNED_날짜시간, 작업자는 UNKNOWN으로 저장될 수 있으므로 검사 전에 반드시 입력하십시오.", "warning")
    add_heading(doc, "설비 상태 확인", 2)
    add_table(doc, ["확인 항목", "정상 기준"], [
        ["PLC 연결", "READY 또는 정상 연결 상태"],
        ["요청 장치", "기본 D810, START_REQ 1, END_REQ 0"],
        ["하중 장치", "기본 D801에서 연속 2 WORD 읽기"],
        ["현재 하중", "대기 상태에서 현장 조건에 맞는 값 표시"],
        ["그래프 단위", "X축 읽기 횟수, Y축 하중 kgf"],
        ["미저장 수량", "이전 LOT 데이터가 남았으면 먼저 저장 여부 결정"],
    ], widths=[4.5, 12.5])
    add_heading(doc, "하중 값 변환", 2)
    add_body(doc, "프로그램은 D801을 하위 WORD, D802를 상위 WORD로 읽어 32비트 부호 있는 정수 원시값을 만듭니다. 표시 및 저장 하중은 원시값을 하중 변환 배율로 나눈 kgf 값입니다. 기본 배율은 1000입니다.")
    add_callout(doc, "예시", "PLC 원시값이 441이고 변환 배율이 1000이면 화면과 CSV에는 0.441 kgf로 표시됩니다.")

    page_break(doc)
    add_heading(doc, "제5장  검사 실행")
    add_heading(doc, "표준 검사 순서", 2)
    add_steps(doc, [
        "작업정보와 PLC 연결 상태를 확인한 뒤 START 버튼을 누릅니다.",
        "버튼이 STOP으로 바뀌고 프로그램이 PLC의 START_REQ를 기다리는지 확인합니다.",
        "PLC 요청 장치가 START 값으로 바뀌면 그래프가 초기화되고 하중 데이터를 읽기 시작합니다.",
        "START 값이 유지되는 동안 LOAD_RAW를 계속 읽어 그래프와 메모리에 추가합니다.",
        "PLC 요청 장치가 END 값으로 바뀌면 수집을 끝내고 최고점으로 GOOD 또는 NG를 판정합니다.",
        "검사 결과가 검사현황과 LOG에 반영되고 완료 데이터가 해당 설비의 LOT 메모리에 누적되는지 확인합니다.",
        "다음 제품도 같은 방식으로 반복합니다. LOT 작업이 끝나면 LOT END를 누릅니다.",
    ])
    add_heading(doc, "PLC 요청과 프로그램 동작", 2)
    add_table(doc, ["PLC 상태", "기본값", "프로그램 동작"], [
        ["START_REQ", "D810 = 1", "새 검사를 시작하고 D801/D802 하중을 반복해서 읽습니다."],
        ["LOAD_RAW", "D801, D802", "32비트 원시값을 kgf로 변환해 그래프와 현재 검사 메모리에 추가합니다."],
        ["END_REQ", "D810 = 0", "수집을 종료하고 최고점 판정 후 완료 데이터를 LOT 메모리에 저장합니다."],
        ["기타 요청값", "1과 0 이외", "유효하지 않은 요청값으로 기록하며 정상 검사로 완료하지 않습니다."],
    ], widths=[3.3, 3.4, 10.3])
    add_heading(doc, "STOP 사용", 2)
    add_body(doc, "STOP은 현재 진행 중인 미완료 수집을 중단합니다. 이미 END_REQ까지 완료되어 LOT 메모리에 누적된 검사 결과는 유지됩니다. 비상 정지나 잘못된 작업정보를 발견했을 때 사용하고, 원인을 조치한 뒤 다시 START하십시오.")
    add_callout(doc, "시뮬레이션 모드", "설정에서 Runtime.Simulation이 1이면 START 후 PLC 대신 임의 하중 데이터 60개를 약 300 ms 간격으로 만듭니다. 시뮬레이션 결과는 실제 측정값이 아니므로 생산 운영 전에 반드시 0인지 확인하십시오.", "warning")

    page_break(doc)
    add_heading(doc, "제6장  그래프와 판정 확인")
    add_heading(doc, "실시간 그래프", 2)
    add_table(doc, ["항목", "표시 기준"], [
        ["X축", "읽기 횟수입니다. 초기 최대값은 설정값에 10% 여유를 더하고 10칸으로 나누어 표시합니다. 데이터가 범위를 넘으면 자동 확장됩니다."],
        ["Y축", "하중 kgf입니다. 검사 SPEC과 설정 범위를 기준으로 10칸을 표시하며 최고 하중이 범위를 넘으면 자동 확장됩니다."],
        ["현재 하중", "하중 장치와 WORD 값, 변환된 kgf를 소수점 셋째 자리까지 표시합니다."],
        ["완료 표시", "검사 종료 시 최고점을 반투명 빨간 원으로 표시하고 최대 하중을 kgf로 적습니다."],
    ], widths=[3.5, 13.5])
    add_heading(doc, "판정 기준", 2)
    add_table(doc, ["판정", "조건", "운영 조치"], [
        ["GOOD", "최고 하중이 SPEC 하한값 이상", "정상적으로 다음 검사를 진행합니다."],
        ["NG", "최고 하중이 SPEC 하한값 미만", "제품을 구분하고 현장 불량 처리 절차를 따릅니다."],
        ["INVALID", "수집 데이터가 없거나 SPEC을 판정할 수 없음", "PLC 요청 순서와 하중 입력, SPEC 설정을 확인하고 재검사합니다."],
    ], widths=[2.7, 6.2, 8.1])
    add_callout(doc, "판정 확인", "그래프 끝의 문자 대신 검사현황 판정, 최고점 표시, LOG 메시지를 함께 확인하십시오. 판정용 값은 검사 구간 전체에서 가장 큰 하중입니다.")

    page_break(doc)
    add_heading(doc, "제7장  LOT 종료와 CSV 저장")
    add_heading(doc, "LOT END 저장 순서", 2)
    add_steps(doc, [
        "해당 설비가 END_REQ까지 완료되어 수집 중이 아닌지 확인합니다.",
        "검사현황의 검사수량, GOOD, NG, Yield(%)를 확인합니다.",
        "LOT END를 누릅니다.",
        "저장 완료 메시지와 LOG를 확인합니다.",
        "저장된 파일을 확인한 뒤 다음 LOT 작업정보를 입력합니다.",
    ])
    add_body(doc, "LOT END는 완료된 검사 데이터가 한 건 이상 있을 때만 실행됩니다. 저장이 성공하면 메모리 데이터가 초기화됩니다. 저장에 실패하면 메모리 데이터가 유지되므로 프로그램을 종료하지 말고 원인을 조치한 뒤 다시 저장하십시오.")
    add_heading(doc, "자동 저장", 2)
    add_body(doc, "설비별 누적 검사수가 LOT 자동 저장 검사 횟수에 도달하면 프로그램이 자동으로 HDD 저장을 실행합니다. 기본값은 600건이며 SETTING에서 설비별로 지정할 수 있습니다.")
    add_heading(doc, "저장 위치와 파일명", 2)
    add_table(doc, ["항목", "형식 또는 예시"], [
        ["저장 폴더", "C:\\SOFCMeas\\Data\\PLC1 또는 PLC2\\yyyy\\MM\\dd"],
        ["파일명", "yyyyMMdd_HHmmss_LOTNO_MODEL.csv"],
        ["예시", "20260906_174418_260827-A01_SOFC_A12.csv"],
        ["중복 파일명", "같은 이름이 있으면 기존 파일을 덮어쓰지 않고 구분 가능한 접미사를 붙입니다."],
    ], widths=[4.0, 13.0])
    add_callout(doc, "주의", "수집 중이거나 저장 중일 때 프로그램을 종료하지 마십시오. 종료 확인 창에 미저장 건수가 표시되면 아니요를 선택하고 LOT END로 먼저 저장하십시오.", "warning")

    page_break(doc)
    add_heading(doc, "제8장  CSV 파일 구성")
    add_body(doc, "CSV는 UTF-8 형식이며 작업정보 두 줄, 빈 줄, 검사 데이터 표 순서로 저장됩니다. 원시 PLC 정수값이 아니라 변환된 kgf 값을 소수점 셋째 자리까지 저장합니다.")
    add_heading(doc, "작업정보 영역", 2)
    add_table(doc, ["행", "저장 내용"], [
        ["1행 항목", "모델명, LOT NO., 작업자, SPEC, 생산수량, GOOD, NG, Yield(%)"],
        ["2행 값", "각 항목에 대응하는 실제 값"],
        ["3행", "빈 줄"],
    ], widths=[3.0, 14.0])
    add_heading(doc, "검사 데이터 영역", 2)
    add_table(doc, ["열", "내용"], [
        ["NO", "LOT 내 검사 순번"],
        ["판정결과", "GOOD, NG 또는 INVALID"],
        ["RAWDATA개수", "해당 검사의 하중 샘플 수"],
        ["1, 2, 3 ...", "각 검사의 샘플 순번에 따른 kgf 값"],
    ], widths=[4.0, 13.0])
    add_body(doc, "샘플 번호 열은 LOT에 포함된 검사 중 RAWDATA개수가 가장 많은 검사를 기준으로 생성됩니다. 샘플 수가 적은 검사 행은 뒤쪽 열을 빈 값으로 둡니다.")
    add_heading(doc, "간단한 형식 예", 2)
    add_table(doc, ["NO", "판정결과", "RAWDATA개수", "1", "2", "3", "4"], [
        ["1", "GOOD", "4", "0.102", "0.254", "0.441", "0.405"],
        ["2", "NG", "3", "0.095", "0.210", "0.318", ""],
    ], widths=[1.2, 2.6, 2.8, 2.6, 2.6, 2.6, 2.6], font_size=8.2)

    page_break(doc)
    add_heading(doc, "제9장  저장 결과 조회와 REVIEW")
    add_image(doc, REVIEW_SCREEN, Inches(4.45), "그림 2  저장 결과 조회와 검사 번호 REVIEW 예시")
    add_steps(doc, [
        "조회할 설비의 저장 결과 조회 영역에서 시작일과 종료일을 선택합니다.",
        "FIND를 눌러 해당 날짜 범위의 CSV 파일 목록을 불러옵니다.",
        "목록에서 확인할 파일 행을 클릭합니다. 프로그램이 CSV를 읽으면 검사 번호 콤보박스가 갱신됩니다.",
        "검사 번호를 선택하고 REVIEW를 누릅니다.",
        "그래프, 최고점 원과 최대 하중, 판정, 검사수량을 확인합니다.",
    ])
    add_callout(doc, "목록 선택", "현재 프로그램은 결과 행 클릭을 기준으로 파일을 읽습니다. 더블클릭이 아니라 행을 한 번 선택한 뒤 검사 번호가 바뀌었는지 확인하십시오.")
    add_heading(doc, "조회 시 확인 사항", 2)
    add_bullets(doc, [
        "설비 1 화면은 PLC1 폴더, 설비 2 화면은 PLC2 폴더의 파일을 조회합니다.",
        "REVIEW 그래프의 X축과 Y축도 각각 10칸으로 표시되고 데이터가 범위를 넘으면 확장됩니다.",
        "최고점은 반투명 빨간 원과 ‘최대 0.000 kgf’ 형식으로 표시됩니다.",
        "구형 또는 손상된 CSV는 파싱 오류가 날 수 있으므로 LOG에서 파일명과 오류 내용을 확인합니다.",
    ])

    page_break(doc)
    add_heading(doc, "제10장  LOG VIEW와 로그 파일")
    add_heading(doc, "LOG VIEW 사용", 2)
    add_body(doc, "LOG VIEW는 ADMINISTRATOR 권한에서 사용할 수 있습니다. 날짜와 설비 전체, 설비 1 또는 설비 2를 선택해 조회하며, 필요한 결과는 CSV로 내보낼 수 있습니다.")
    add_steps(doc, [
        "상단 USER 표시를 눌러 ADMINISTRATOR로 전환합니다.",
        "LOG VIEW를 누릅니다.",
        "조회 날짜와 설비를 선택한 뒤 조회합니다.",
        "시간, 설비, LOT, 파일, 판정, 최고점, 메시지를 확인합니다.",
        "현장 분석에 필요한 경우 CSV 내보내기를 실행합니다.",
    ])
    add_heading(doc, "응용 프로그램 로그", 2)
    add_table(doc, ["항목", "내용"], [
        ["기본 위치", "C:\\SOFCMeas\\Log\\yyyy\\MM\\dd"],
        ["파일명", "SOFCMeas_yyyyMMdd_001.log부터 순차 생성"],
        ["기록 내용", "시각, 세션, 순번, PID, TID, 레벨, 발생 위치, 이벤트"],
        ["주요 이벤트", "실행과 권한, PLC 연결과 재연결, 요청값 변화, 검사 시작과 종료, 샘플, 판정, 메모리 누적, LOT 저장, 오류"],
        ["파일 분할", "기본 최대 20 MB에 도달하면 새 로그 파일 생성"],
        ["보존 정책", "기본 30일, 자동 삭제 사용"],
    ], widths=[4.0, 13.0], font_size=8.4)
    add_callout(doc, "장애 분석", "문제가 발생하면 발생 시각, 설비 번호, LOT NO., 화면 메시지를 기록하고 해당 날짜의 전체 로그 폴더를 보존하십시오. 로그만으로 요청값, 하중 읽기, 판정, 저장 순서를 추적할 수 있습니다.")

    page_break(doc)
    add_heading(doc, "제11장  SETTING")
    add_image(doc, SETTING_SCREEN, Inches(6.6), "그림 3  공통 설정과 설비 1, 설비 2 설정 화면 예시이며 현장 값은 다를 수 있음")
    add_body(doc, "SETTING은 ADMINISTRATOR만 저장할 수 있습니다. 시뮬레이션 설정은 다음 START부터 적용되며, 그 밖의 설정은 프로그램을 재시작한 뒤 적용됩니다.")
    add_heading(doc, "공통 설정", 2)
    add_table(doc, ["항목", "기본값", "설명"], [
        ["Runtime.Simulation", "0", "0은 PLC 운영, 1은 시뮬레이션"],
        ["Security.AdminPassword", "1234", "관리자 전환 암호. 현장 적용 전에 변경"],
        ["Log.RetentionDays", "30", "로그 보존일"],
        ["Log.AutoDelete", "true", "보존일이 지난 로그 자동 삭제"],
        ["Log.MaximumFileSizeMb", "20", "로그 파일 한 개의 최대 크기"],
    ], widths=[6.0, 2.8, 8.2], font_size=8.3)
    add_callout(doc, "설정 변경 원칙", "생산 중에는 설정을 바꾸지 마십시오. 변경 전에 현재 값을 기록하고, SAVE 후 프로그램을 재시작한 다음 설비 1과 설비 2를 각각 시험하십시오.", "warning")

    page_break(doc)
    add_heading(doc, "제12장  설비별 PLC와 검사 설정")
    add_table(doc, ["항목", "기본값", "운영 의미"], [
        ["PLC IP", "#1 172.20.9.100 / #2 172.20.9.101", "각 설비 PLC 주소"],
        ["PLC 자동 연결", "true", "시작 시 자동 연결"],
        ["PLC 포트", "10000", "PLC TCP 통신 포트"],
        ["읽기 주기", "1000 ms", "PLC 요청과 하중을 확인하는 주기"],
        ["재연결 간격", "2000 ms", "통신 실패 후 재시도 간격"],
        ["통신 Timeout", "3000 ms", "응답 대기 한도"],
        ["하중 변환 배율", "1000", "원시값을 kgf로 나누는 값"],
        ["요청 디바이스", "D810", "START_REQ와 END_REQ 장치"],
        ["하중 디바이스", "D801", "연속 2 WORD 하중 장치"],
        ["START 요청값", "1", "검사 수집 시작"],
        ["END 요청값", "0", "검사 수집 종료와 판정"],
        ["LOT 자동 저장 검사 횟수", "600", "메모리 누적 후 자동 저장 건수"],
        ["검사 SPEC", "0.40 kgf", "GOOD와 NG를 나누는 최고점 하한"],
        ["X축 자동 확장", "true", "읽기 횟수가 표시 범위를 넘으면 확장"],
        ["Y축 자동 확장", "true", "하중이 표시 범위를 넘으면 확장"],
    ], widths=[5.3, 4.7, 7.0], font_size=7.7)
    add_heading(doc, "통신 설정 변경 후 확인", 2)
    add_steps(doc, [
        "설비 번호와 IP가 서로 바뀌지 않았는지 확인합니다.",
        "요청 장치, 하중 장치, START와 END 값이 PLC 시퀀스와 같은지 확인합니다.",
        "하중 변환 배율을 변경했다면 기준 하중으로 kgf 표시를 검증합니다.",
        "SAVE 후 프로그램을 재시작하고 각 설비에서 한 건씩 시험합니다.",
        "시험 결과를 LOT END로 저장하고 CSV 값과 REVIEW 그래프를 확인합니다.",
    ])

    page_break(doc)
    add_heading(doc, "제13장  정상 종료")
    add_steps(doc, [
        "설비 1과 설비 2 모두 검사 수집 중이 아닌지 확인합니다.",
        "각 설비의 미저장 검사수량이 있으면 LOT END로 저장합니다.",
        "저장 완료 메시지와 생성된 CSV 파일을 확인합니다.",
        "화면 오른쪽 위 종료 버튼을 누릅니다.",
        "‘프로그램을 종료하시겠습니까?’에서 예를 선택합니다.",
    ])
    add_body(doc, "수집 중이거나 저장하지 않은 데이터가 있으면 종료 확인 창에 HDD 미저장 건수가 표시됩니다. 종료하면 메모리 데이터가 사라지므로 아니요를 선택하고 저장한 뒤 다시 종료하십시오.")
    add_callout(doc, "강제 종료 금지", "LOT 저장 중에는 Windows 종료, 전원 차단, 작업 관리자 강제 종료를 하지 마십시오. CSV 파일이 완전하게 생성되지 않을 수 있습니다.", "warning")
    add_heading(doc, "종료 전 체크", 2)
    add_table(doc, ["설비", "수집 상태", "미저장 수량", "CSV 확인", "완료"], [
        ["설비 1", "정지", "0건", "PLC1 날짜 폴더", "□"],
        ["설비 2", "정지", "0건", "PLC2 날짜 폴더", "□"],
    ], widths=[3.0, 3.0, 3.4, 5.0, 2.6])

    page_break(doc)
    add_heading(doc, "제14장  문제 대응")
    add_table(doc, ["증상", "확인 순서"], [
        ["프로그램이 시작 직후 종료됨", "표시된 영어 메시지를 기록합니다. MAC Address 등록, 40일 기간, Windows 날짜와 시간을 확인한 뒤 관리자에게 문의합니다."],
        ["설비가 연결 대기 또는 오류 상태", "PLC 전원과 케이블을 확인합니다. 설비별 IP, 포트, 방화벽을 확인하고 해당 시각의 PLC 연결·재연결 로그를 봅니다."],
        ["START를 눌러도 그래프가 그려지지 않음", "시뮬레이션 여부, D810 요청값, START 값 1, D801/D802 데이터 변화를 확인합니다. LOG의 요청값과 샘플 이벤트를 확인합니다."],
        ["하중 값이 실제와 다름", "D801/D802 WORD 순서와 부호, 하중 변환 배율 1000을 확인합니다. 기준 하중으로 화면과 CSV의 kgf 값을 비교합니다."],
        ["판정이 INVALID", "END_REQ 전에 샘플이 들어왔는지, SPEC 값이 유효한지 확인합니다. 잘못된 검사는 재검사합니다."],
        ["LOT END가 실행되지 않음", "현재 검사가 END_REQ까지 끝났는지와 완료 데이터가 한 건 이상인지 확인합니다. 저장 중이면 완료될 때까지 기다립니다."],
        ["LOT 저장 실패", "C:\\SOFCMeas\\Data 접근 권한, 디스크 여유 공간, 파일 사용 여부를 확인합니다. 메모리가 유지되므로 프로그램을 종료하지 말고 다시 저장합니다."],
        ["저장 결과가 조회되지 않음", "조회 날짜와 설비가 맞는지 확인합니다. PLC1/PLC2 날짜 폴더에 CSV가 있는지 보고 FIND 후 결과 행을 클릭합니다."],
        ["검사 번호가 바뀌지 않음", "결과 행을 한 번 클릭해 CSV 로드가 끝날 때까지 기다립니다. 콤보박스 항목을 확인한 뒤 REVIEW를 누릅니다."],
        ["과거 CSV 파싱 오류", "작업정보 2줄, 빈 줄, NO/판정결과/RAWDATA개수 헤더가 있는지 확인합니다. 파일과 로그를 관리자에게 전달합니다."],
    ], widths=[5.1, 11.9], font_size=8.0)
    add_callout(doc, "관리자 전달 정보", "발생 날짜와 시각, 설비 번호, LOT NO., 화면 메시지, 재현 순서, C:\\SOFCMeas\\Log의 해당 날짜 폴더, 관련 CSV 파일을 함께 전달하십시오.")

    page_break(doc)
    add_heading(doc, "제15장  작업자 점검표")
    add_heading(doc, "작업 시작 전", 2)
    add_table(doc, ["점검 항목", "확인"], [
        ["PC 날짜와 시간 정상", "□"],
        ["설비 1 PLC 연결 정상", "□"],
        ["설비 2 PLC 연결 정상", "□"],
        ["시뮬레이션 모드 0 확인", "□"],
        ["모델명, LOT NO., 작업자 입력", "□"],
        ["SPEC 값 확인", "□"],
        ["이전 LOT 미저장 데이터 없음", "□"],
    ], widths=[14.5, 2.5], font_size=7.8)
    add_heading(doc, "LOT 종료 시", 2)
    add_table(doc, ["점검 항목", "확인"], [
        ["마지막 검사가 END_REQ까지 완료", "□"],
        ["검사수량, GOOD, NG, Yield 확인", "□"],
        ["LOT END 실행", "□"],
        ["저장 완료 메시지 확인", "□"],
        ["PLC1 또는 PLC2 날짜 폴더의 CSV 확인", "□"],
        ["다음 LOT 작업정보 갱신", "□"],
    ], widths=[14.5, 2.5], font_size=7.8)
    add_heading(doc, "핵심 경로", 2)
    add_table(doc, ["용도", "경로"], [
        ["프로그램 기본 폴더", "C:\\SOFCMeas"],
        ["설비 1 데이터", "C:\\SOFCMeas\\Data\\PLC1\\yyyy\\MM\\dd"],
        ["설비 2 데이터", "C:\\SOFCMeas\\Data\\PLC2\\yyyy\\MM\\dd"],
        ["응용 프로그램 로그", "C:\\SOFCMeas\\Log\\yyyy\\MM\\dd"],
        ["설정", "C:\\SOFCMeas\\Conf"],
    ], widths=[5.0, 12.0], font_size=7.8)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    build_document()
